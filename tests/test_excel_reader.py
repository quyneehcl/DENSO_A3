import pandas as pd
import pytest
from openpyxl import Workbook

from ingest.excel_reader import read_excel
from ingest.schema import validate_chunks


@pytest.fixture
def error_code_xlsx(tmp_path):
    """Title row, header, vertically merged component name."""
    wb = Workbook()
    ws = wb.active
    ws.title = "MaLoi"
    ws.append(["BẢNG MÃ LỖI MÁY ÉP"])
    ws.merge_cells("A1:C1")
    ws.append([])
    ws.append(["Tên linh kiện", "Mã lỗi", "Cách xử lý"])
    ws.append(["Cảm biến nhiệt", "E01", "Thay cảm biến"])
    ws.append([None, "E02", "Kiểm tra dây nối"])
    ws.merge_cells("A4:A5")
    ws.append(["Động cơ servo", "E10", "Reset driver, kiểm tra encoder"])
    p = tmp_path / "ma_loi.xlsx"
    wb.save(p)
    return p


def test_merged_cells_and_rows(error_code_xlsx):
    chunks = [c.to_dict() for c in read_excel(error_code_xlsx, doc_id="doc001")]
    assert validate_chunks(chunks) == []
    assert [c["chunk_id"] for c in chunks] == ["doc001_chunk001", "doc001_chunk002", "doc001_chunk003"]
    assert chunks[0]["text"] == "Tên linh kiện là Cảm biến nhiệt, mã lỗi là E01, cách xử lý là Thay cảm biến."
    # vertically merged component propagated to row 5
    assert chunks[1]["text"].startswith("Tên linh kiện là Cảm biến nhiệt, mã lỗi là E02")
    assert chunks[1]["location"] == "MaLoi, hàng 5"
    assert chunks[2]["location"] == "MaLoi, hàng 6"
    assert "Reset driver, kiểm tra encoder" in chunks[2]["text"]
    assert all(c["doc_type"] == "excel_table" and c["language"] == "vi" for c in chunks)
    assert all(c["source_file"] == "ma_loi.xlsx" for c in chunks)


def test_two_level_header_multi_sheet(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "May ep"
    ws.append(["Model", "Thông số điện", None, "Công suất (kW)"])
    ws.append([None, "Điện áp (V)", "Tần số (Hz)", None])
    ws.merge_cells("B1:C1")
    ws.merge_cells("A1:A2")
    ws.merge_cells("D1:D2")
    ws.append(["MX-200", 380, 50, 7.5])
    ws2 = wb.create_sheet("Robot")
    ws2.append(["Model", "Tải trọng (kg)"])
    ws2.append(["RB-10", 10])
    p = tmp_path / "thong_so.xlsx"
    wb.save(p)
    chunks = read_excel(p, doc_id="doc002")
    texts = [c.text for c in chunks]
    assert texts[0] == ("Model là MX-200, thông số điện - Điện áp (V) là 380, "
                        "thông số điện - Tần số (Hz) là 50, công suất (kW) là 7.5.")
    assert chunks[0].location == "May ep, hàng 3"
    assert texts[1] == "Model là RB-10, tải trọng (kg) là 10."
    assert chunks[1].location == "Robot, hàng 2"


def test_pandas_path_keeps_row_numbers(tmp_path):
    p = tmp_path / "simple.xlsx"
    with pd.ExcelWriter(p) as w:
        pd.DataFrame({"Mã": ["A1", "A2"], "Mô tả": ["Lỗi áp suất", None]}).to_excel(
            w, sheet_name="S1", index=False, startrow=2)
    chunks = read_excel(p, doc_id="d")
    assert [c.location for c in chunks] == ["S1, hàng 4", "S1, hàng 5"]
    assert chunks[1].text == "Mã là A2."
