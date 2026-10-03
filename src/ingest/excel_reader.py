"""Excel -> intermediate chunks (doc_type = "excel_table").

Each data row becomes one structured sentence (one chunk), e.g.
"Tên linh kiện là Cảm biến nhiệt, mã lỗi là E01, cách xử lý là Thay cảm biến."

Reading strategy:
  * pandas (all sheets) for simple sheets;
  * openpyxl raw read when a sheet has merged cells: every cell of a merged range
    receives the top-left value (handles merged group headers and vertically merged
    values such as one component spanning several error codes).
"""
from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

import pandas as pd

from .schema import Chunk, default_doc_id, detect_language, make_chunk_id
from .table_utils import format_cell, row_to_sentence, split_header

log = logging.getLogger(__name__)

DOC_TYPE = "excel_table"


@dataclass
class SheetGrid:
    name: str
    rows: list[tuple[int, list[str]]]  # (excel row number, cell texts)


def _merged_sheet_names(path: Path) -> set[str]:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return set()
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=False)
    try:
        return {ws.title for ws in wb.worksheets if ws.merged_cells.ranges}
    finally:
        wb.close()


def _read_with_openpyxl(path: Path, sheet: str) -> SheetGrid:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    try:
        ws = wb[sheet]
        grid = [[c.value for c in row] for row in ws.iter_rows(min_row=1, min_col=1)]
        for rng in ws.merged_cells.ranges:
            top_left = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if r - 1 < len(grid) and c - 1 < len(grid[r - 1]):
                        grid[r - 1][c - 1] = top_left
    finally:
        wb.close()
    return SheetGrid(sheet, [(i + 1, [format_cell(v) for v in row]) for i, row in enumerate(grid)])


def _read_with_pandas(path: Path, sheet) -> SheetGrid:
    df = pd.read_excel(path, sheet_name=sheet, header=None, dtype=object)
    return SheetGrid(
        str(sheet),
        # header=None keeps every row from row 1, so index i <-> Excel row i + 1
        [(int(i) + 1, [format_cell(v) for v in row]) for i, row in zip(df.index, df.itertuples(index=False))],
    )


def load_sheets(path: str | Path) -> list[SheetGrid]:
    path = Path(path)
    if path.suffix.lower() == ".csv":
        df = pd.read_csv(path, header=None, dtype=object)
        return [SheetGrid(path.stem, [(i + 1, [format_cell(v) for v in r]) for i, r in enumerate(df.itertuples(index=False))])]
    merged = _merged_sheet_names(path)
    sheets = []
    for name in pd.ExcelFile(path).sheet_names:
        if name in merged:
            log.debug("%s / %s: merged cells -> openpyxl", path.name, name)
            sheets.append(_read_with_openpyxl(path, name))
        else:
            sheets.append(_read_with_pandas(path, name))
    return sheets


def _is_title_row(cells: list[str]) -> bool:
    """Title / caption row: a single distinct value (a merged title fills several cells)."""
    filled = {c for c in cells if c}
    return len(filled) == 1


@dataclass
class TableRow:
    sheet: str
    row_number: int
    headers: list[str]
    values: list[str]


def extract_table_rows(sheet: SheetGrid) -> list[TableRow]:
    """Segment a sheet into tables (blocks separated by empty rows) and return data rows.

    A block starting with a title row, or the first block of the sheet, opens a new table
    whose header is detected automatically. Other blocks continue the previous table.
    """
    blocks: list[list[tuple[int, list[str]]]] = []
    current: list[tuple[int, list[str]]] = []
    for num, cells in sheet.rows:
        if any(cells):
            current.append((num, cells))
        elif current:
            blocks.append(current)
            current = []
    if current:
        blocks.append(current)

    out: list[TableRow] = []
    headers: list[str] | None = None
    header_raw: list[list[str]] = []
    for block in blocks:
        i = 0
        had_title = False
        while i < len(block) and _is_title_row(block[i][1]):
            had_title = True
            i += 1
        if i >= len(block):
            continue  # titles / notes only
        if headers is None or had_title:
            raw = [cells for _, cells in block[i:]]
            headers, n = split_header(raw)
            header_raw = raw[:n]
            i += n
        for num, cells in block[i:]:
            if cells in header_raw:  # header repeated (e.g. printed on every page)
                continue
            out.append(TableRow(sheet.name, num, headers, cells))
    return out


def read_excel(path: str | Path, doc_id: str | None = None, start_index: int = 1) -> list[Chunk]:
    path = Path(path)
    doc_id = doc_id or default_doc_id(path)
    chunks: list[Chunk] = []
    for sheet in load_sheets(path):
        rows = extract_table_rows(sheet)
        for tr in rows:
            lang = detect_language(" ".join(tr.headers + tr.values))
            text = row_to_sentence(tr.headers, tr.values, lang)
            if not text:
                continue
            chunks.append(
                Chunk(
                    chunk_id=make_chunk_id(doc_id, start_index + len(chunks)),
                    text=text,
                    source_file=path.name,
                    location=f"{sheet.name}, hàng {tr.row_number}",
                    doc_type=DOC_TYPE,
                    language=lang,
                )
            )
    return chunks
